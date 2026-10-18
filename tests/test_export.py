"""Export pre účtovníka: obdobia, výber dokladov, CSV, Excel, ZIP s PDF."""

import io
import os
import sys
import zipfile
from datetime import date, timedelta
from decimal import Decimal
from pathlib import Path

os.environ.setdefault("SECRET_KEY", "test-secret-key")
os.environ.setdefault("SESSION_HTTPS_ONLY", "false")

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook
from sqlalchemy import create_engine
from sqlalchemy.orm import sessionmaker
from sqlalchemy.pool import StaticPool

from audit_log import AuditLog
from auth import require_login_api, require_login_page
from csrf import verify_csrf
from database import Base, get_db
from export_utils import (
    INVOICE_COLUMNS,
    build_csv,
    resolve_period,
    safe_text,
)
from main import app
from models import Company, Customer, Invoice, InvoiceItem, Job, JobCost

test_engine = create_engine(
    "sqlite://", connect_args={"check_same_thread": False}, poolclass=StaticPool
)

TestingSessionLocal = sessionmaker(autocommit=False, autoflush=False, bind=test_engine)

# pevný dátum v strede mesiaca - nezávislé od dňa spustenia testov
D = date(2026, 3, 15)
_counter = {"n": 0}


@pytest.fixture(autouse=True)
def setup_test_database():

    _counter["n"] = 0

    Base.metadata.drop_all(bind=test_engine)
    Base.metadata.create_all(bind=test_engine)

    db = TestingSessionLocal()
    db.add(Customer(name="Stavebniny s.r.o.", ico="12345008", dic="2020123456", ic_dph="SK2020123456"))
    db.add(Customer(name="=HYPERLINK(\"http://zle.sk\")"))
    db.add(Company(id=1, name="Moja Firma", ico="12345016", is_vat_payer=True))
    db.add(Job(title="Oprava strechy", customer_id=1))
    db.commit()
    db.close()

    def override_get_db():

        db = TestingSessionLocal()

        try:
            yield db

        finally:
            db.close()

    app.dependency_overrides[get_db] = override_get_db
    app.dependency_overrides[require_login_page] = lambda: "testuser"
    app.dependency_overrides[require_login_api] = lambda: "testuser"
    app.dependency_overrides[verify_csrf] = lambda: None

    yield

    app.dependency_overrides.clear()


client = TestClient(app)

MARCH = "preset=custom&date_from=2026-03-01&date_to=2026-03-31"


def add_invoice(issue=D, status="Odoslaná", price="1000", vat=23, customer=1,
                credit_note=False, proforma=False, original=None, paid=None):

    _counter["n"] += 1

    db = TestingSessionLocal()

    invoice = Invoice(
        invoice_number=f"2026{_counter['n']:03d}",
        customer_id=customer,
        status=status,
        issue_date=issue,
        due_date=issue + timedelta(days=14),
        paid_date=paid,
        is_credit_note=credit_note,
        is_proforma=proforma,
        original_invoice_id=original,
    )

    db.add(invoice)
    db.flush()

    db.add(InvoiceItem(
        invoice_id=invoice.id, description="Práca", quantity=Decimal("1"),
        unit="ks", unit_price=Decimal(price), vat_rate=vat,
    ))

    db.commit()
    invoice_id = invoice.id
    db.close()

    return invoice_id


def add_cost(amount="100", cost_date=D, category="material", hours=None):

    db = TestingSessionLocal()

    db.add(JobCost(
        job_id=1, description="Náklad", amount=Decimal(amount), cost_date=cost_date,
        category=category, hours=Decimal(hours) if hours else None,
    ))

    db.commit()
    db.close()


# ---------- obdobie ----------

def test_periods_around_year_boundary():

    today = date(2026, 1, 15)

    assert resolve_period("last_month", None, None, today)[:2] == (date(2025, 12, 1), date(2025, 12, 31))
    assert resolve_period("this_month", None, None, today)[:2] == (date(2026, 1, 1), date(2026, 1, 31))
    assert resolve_period("last_quarter", None, None, today)[:2] == (date(2025, 10, 1), date(2025, 12, 31))
    assert resolve_period("this_quarter", None, None, today)[:2] == (date(2026, 1, 1), date(2026, 3, 31))
    assert resolve_period("last_year", None, None, today)[:2] == (date(2025, 1, 1), date(2025, 12, 31))


def test_february_and_leap_year_month_end():

    assert resolve_period("this_month", None, None, date(2028, 2, 10))[1] == date(2028, 2, 29)
    assert resolve_period("this_month", None, None, date(2026, 2, 10))[1] == date(2026, 2, 28)


def test_custom_period_validation():

    assert resolve_period("custom", date(2026, 1, 1), date(2026, 1, 31), D)[:2] == (date(2026, 1, 1), date(2026, 1, 31))

    for args in ((None, None), (date(2026, 2, 1), date(2026, 1, 1)), (date(2000, 1, 1), date(2026, 1, 1))):

        with pytest.raises(ValueError):
            resolve_period("custom", *args, D)

    with pytest.raises(ValueError):
        resolve_period("neznáme", None, None, D)


# ---------- bezpečnosť textov ----------

def test_safe_text_neutralizes_formulas():

    for dangerous in ("=1+1", "+cmd", "-2", "@SUM(A1)", "\tx"):
        assert safe_text(dangerous).startswith("'")

    assert safe_text("Normálny text") == "Normálny text"
    assert safe_text(None) == ""


# ---------- CSV ----------

def csv_text(path):

    response = client.get(path)

    assert response.status_code == 200

    return response.content, response


def test_invoices_csv_format():

    add_invoice(price="1000", vat=23)

    content, response = csv_text(f"/export/faktury.csv?{MARCH}")

    assert content.startswith(b"\xef\xbb\xbf")                       # BOM pre Excel
    assert response.headers["content-type"].startswith("text/csv")
    assert 'filename="faktury-2026-03-01_2026-03-31.csv"' in response.headers["content-disposition"]

    lines = content.decode("utf-8-sig").split("\r\n")

    assert lines[0] == ";".join(INVOICE_COLUMNS)

    row = lines[1].split(";")

    assert row[0] == "2026001"
    assert row[2] == "15.03.2026"
    assert row[10:13] == ["1000,00", "230,00", "1230,00"]          # desatinná čiarka


def test_csv_neutralizes_formula_in_customer_name():

    add_invoice(customer=2)

    content, _ = csv_text(f"/export/faktury.csv?{MARCH}")

    assert "'=HYPERLINK" in content.decode("utf-8-sig")


# ---------- výber dokladov ----------

def numbers(path):

    content, _ = csv_text(path)

    return [line.split(";")[0] for line in content.decode("utf-8-sig").split("\r\n")[1:] if line]


def test_selection_by_issue_date_and_status():

    inside = add_invoice(D)
    add_invoice(date(2026, 2, 28))                         # mimo obdobia
    add_invoice(date(2026, 4, 1))                          # mimo obdobia
    add_invoice(D, status="Návrh")                         # návrh sa nezahŕňa
    add_invoice(D, status="Uhradená", paid=D)
    add_invoice(D, status="Stornovaná")
    add_invoice(D, proforma=True)

    result = numbers(f"/export/faktury.csv?{MARCH}")

    assert len(result) == 3
    assert result[0] == "2026001"

    with_proforma = numbers(f"/export/faktury.csv?{MARCH}&proforma=1")

    assert len(with_proforma) == 4


def test_period_boundaries_are_inclusive():

    add_invoice(date(2026, 3, 1))
    add_invoice(date(2026, 3, 31))

    assert len(numbers(f"/export/faktury.csv?{MARCH}")) == 2


def test_credit_note_negative_and_references_original_outside_period():

    original = add_invoice(date(2026, 2, 10), price="500", vat=0)
    add_invoice(D, credit_note=True, price="200", vat=0, original=original)

    content, _ = csv_text(f"/export/faktury.csv?{MARCH}")

    row = content.decode("utf-8-sig").split("\r\n")[1].split(";")

    assert row[1] == "Dobropis"
    assert row[10] == "-200,00"
    assert row[17] == "2026001"                                  # K faktúre


def test_cancelled_listed_but_excluded_from_totals():

    add_invoice(price="1000", vat=0)
    add_invoice(status="Stornovaná", price="9999", vat=0)

    page = client.get(f"/export?{MARCH}").text

    assert "1000.00 €" in page
    assert "9999" not in page
    assert "1 stornovaných" in page

    content, _ = csv_text(f"/export/faktury.csv?{MARCH}")

    assert content.decode("utf-8-sig").count("Stornovaná") == 1


# ---------- náklady ----------

def test_costs_csv_filtered_by_period():

    add_cost("120.50", D, "material")
    add_cost("80", D, "praca", hours="4")
    add_cost("999", date(2026, 5, 1))

    content, _ = csv_text(f"/export/naklady.csv?{MARCH}")

    lines = content.decode("utf-8-sig").split("\r\n")

    assert len([line for line in lines[1:] if line]) == 2
    assert "120,50" in content.decode("utf-8-sig")
    assert "Práca" in content.decode("utf-8-sig")
    assert "999" not in content.decode("utf-8-sig")


# ---------- Excel ----------

def workbook(path):

    response = client.get(path)

    assert response.status_code == 200
    assert "spreadsheetml" in response.headers["content-type"]

    return load_workbook(io.BytesIO(response.content)), response


def test_xlsx_structure_and_types():

    add_invoice(price="1000", vat=23, status="Uhradená", paid=D)
    add_invoice(price="500", vat=0)
    add_cost("300")

    wb, response = workbook(f"/export/uctovnik.xlsx?{MARCH}")

    assert 'filename="uctovnik-2026-03-01_2026-03-31.xlsx"' in response.headers["content-disposition"]
    assert wb.sheetnames == ["Súhrn", "Faktúry", "Rekapitulácia DPH", "Náklady"]

    sheet = wb["Faktúry"]

    assert [cell.value for cell in sheet[1]] == list(INVOICE_COLUMNS)

    first = {sheet.cell(1, c).value: sheet.cell(2, c).value for c in range(1, len(INVOICE_COLUMNS) + 1)}

    assert isinstance(first["Základ"], (int, float)) and first["Základ"] == 1000
    assert first["Spolu"] == 1230.0
    assert first["Dátum vystavenia"].date() == D
    assert first["Dátum úhrady"].date() == D
    assert sheet.cell(2, 11).number_format == "#,##0.00"
    assert sheet.freeze_panes == "A2"


def test_xlsx_summary_and_vat_recap():

    add_invoice(price="1000", vat=23)
    add_invoice(price="500", vat=0)
    add_invoice(price="100", vat=23, credit_note=True)
    add_cost("300")

    wb, _ = workbook(f"/export/uctovnik.xlsx?{MARCH}")

    summary = {row[0].value: row[1].value for row in wb["Súhrn"].iter_rows() if row[0].value}

    assert summary["Obdobie"] == "01.03.2026 – 31.03.2026"
    assert summary["Firma"] == "Moja Firma"
    assert summary["Základ dane"] == 1400.0                       # 1000 + 500 - 100
    assert summary["DPH"] == 207.0                                 # 230 - 23
    assert summary["Spolu s DPH"] == 1607.0
    assert summary["Náklady zo zákaziek (bez DPH)"] == 300.0
    assert summary["Základ − náklady"] == 1100.0

    recap = {row[0].value: [c.value for c in row[1:]] for row in wb["Rekapitulácia DPH"].iter_rows(min_row=2)}

    assert recap["23 %"] == [900.0, 207.0, 1107.0]
    assert recap["0 %"] == [500.0, 0.0, 500.0]
    assert recap["Spolu"] == [1400.0, 207.0, 1607.0]


def test_xlsx_formula_injection_is_stored_as_text():

    add_invoice(customer=2)

    wb, _ = workbook(f"/export/uctovnik.xlsx?{MARCH}")

    cell = wb["Faktúry"].cell(2, 7)

    assert cell.data_type != "f"
    assert str(cell.value).startswith("'=")


def test_xlsx_for_empty_period_is_still_valid():

    wb, _ = workbook(f"/export/uctovnik.xlsx?{MARCH}")

    assert wb["Faktúry"].max_row == 1


# ---------- ZIP s PDF ----------

def test_pdf_zip_contains_one_pdf_per_invoice():

    add_invoice()
    add_invoice()

    response = client.get(f"/export/faktury-pdf.zip?{MARCH}")

    assert response.status_code == 200
    assert response.headers["content-type"] == "application/zip"

    with zipfile.ZipFile(io.BytesIO(response.content)) as archive:

        assert sorted(archive.namelist()) == ["faktura-2026001.pdf", "faktura-2026002.pdf"]
        assert archive.read("faktura-2026001.pdf").startswith(b"%PDF")


def test_pdf_zip_empty_period_is_404():

    assert client.get(f"/export/faktury-pdf.zip?{MARCH}").status_code == 404


# ---------- stránka a validácia ----------

def test_page_shows_preview_and_download_links():

    add_invoice(price="1000", vat=23)
    add_cost("300")

    page = client.get(f"/export?{MARCH}").text

    assert "1000.00 €" in page
    assert "230.00 €" in page
    assert "300.00 €" in page
    assert "/export/uctovnik.xlsx?preset=custom" in page
    assert "/export/faktury.csv?" in page
    assert "/export/faktury-pdf.zip?" in page


def test_page_empty_state_and_default_period_renders():

    assert "V tomto období nič nie je" in client.get(f"/export?{MARCH}").text
    assert client.get("/export").status_code == 200


def test_invalid_dates_are_rejected():

    assert "Neplatný dátum" in client.get("/export?preset=custom&date_from=zle&date_to=2026-01-01").text

    for path in ("faktury.csv", "naklady.csv", "uctovnik.xlsx", "faktury-pdf.zip"):

        assert client.get(f"/export/{path}?preset=custom&date_from=2026-02-01&date_to=2026-01-01").status_code == 422
        assert client.get(f"/export/{path}?preset=custom&date_from=zle&date_to=2026-01-01").status_code == 422
        assert client.get(f"/export/{path}?preset=nezname").status_code == 422


def test_downloads_are_audited():

    add_invoice()

    client.get(f"/export/uctovnik.xlsx?{MARCH}")
    client.get(f"/export/faktury.csv?{MARCH}")

    db = TestingSessionLocal()
    entries = [e.detail for e in db.query(AuditLog).filter(AuditLog.action == "export.download")]
    db.close()

    assert any("Excel" in entry for entry in entries)
    assert any("CSV faktúr" in entry for entry in entries)


def test_export_requires_login():

    app.dependency_overrides.pop(require_login_page, None)

    anonymous = TestClient(app)

    for path in ("/export", "/export/uctovnik.xlsx", "/export/faktury.csv"):
        assert anonymous.get(path, follow_redirects=False).status_code == 303
