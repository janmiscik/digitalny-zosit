"""
Export pre účtovníka: výber dokladov a nákladov za obdobie a ich zápis
do CSV a Excelu. Čistá logika bez HTTP.

Pravidlá:
  * Faktúry sa vyberajú podľa DÁTUMU VYSTAVENIA. Zahrnuté sú vystavené
    doklady (Odoslaná, Uhradená, Stornovaná) - návrhy nie, zálohové len
    na požiadanie. Stornované sú v zozname vidieť, ale do súčtov a
    rekapitulácie DPH sa nepočítajú.
  * Dobropisy majú sumy so záporným znamienkom.
  * Texty začínajúce znakmi =, +, -, @ sa pri zápise neutralizujú
    (ochrana pred vložením vzorca do Excelu cez meno zákazníka).
"""

import csv
import io
from datetime import date, timedelta
from decimal import Decimal

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from invoice_utils import calculate_invoice_totals

ZERO = Decimal("0")

EXPORT_STATUSES = ("Odoslaná", "Uhradená", "Stornovaná")

PRESETS = (
    ("last_month", "Minulý mesiac"),
    ("this_month", "Tento mesiac"),
    ("last_quarter", "Minulý štvrťrok"),
    ("this_quarter", "Tento štvrťrok"),
    ("last_year", "Minulý rok"),
    ("this_year", "Tento rok"),
    ("custom", "Vlastné obdobie"),
)

MAX_PERIOD_DAYS = 366 * 5

INVOICE_COLUMNS = (
    "Číslo", "Typ", "Dátum vystavenia", "Dátum dodania", "Dátum splatnosti",
    "Dátum úhrady", "Odberateľ", "IČO", "DIČ", "IČ DPH", "Základ", "DPH",
    "Spolu", "Stav", "Spôsob úhrady", "Variabilný symbol",
    "Prenesená daňová povinnosť", "K faktúre", "V súčte",
)

COST_COLUMNS = ("Dátum", "Zákazka", "Zákazník", "Druh", "Popis", "Hodiny", "Suma bez DPH")

COST_LABELS = {"material": "Materiál", "praca": "Práca", "ine": "Iné"}


# =========================================
# OBDOBIE
# =========================================

def _month_start(year: int, month: int) -> date:
    return date(year, month, 1)


def _month_end(year: int, month: int) -> date:

    if month == 12:
        return date(year, 12, 31)

    return date(year, month + 1, 1) - timedelta(days=1)


def resolve_period(preset: str, date_from: date | None, date_to: date | None,
                   today: date) -> tuple[date, date, str]:
    """Vráti (od, do, popis). Vyhodí ValueError pri neplatnom vstupe."""

    if preset == "custom":

        if date_from is None or date_to is None:
            raise ValueError("Pri vlastnom období zadaj dátum od aj do.")

        if date_from > date_to:
            raise ValueError("Dátum 'od' nesmie byť po dátume 'do'.")

        if (date_to - date_from).days > MAX_PERIOD_DAYS:
            raise ValueError("Obdobie môže mať najviac 5 rokov.")

        return date_from, date_to, f"{date_from:%d.%m.%Y} – {date_to:%d.%m.%Y}"

    if preset == "this_month":
        return _month_start(today.year, today.month), _month_end(today.year, today.month), f"{today.month:02d}/{today.year}"

    if preset == "last_month":

        year, month = (today.year, today.month - 1) if today.month > 1 else (today.year - 1, 12)

        return _month_start(year, month), _month_end(year, month), f"{month:02d}/{year}"

    if preset in ("this_quarter", "last_quarter"):

        quarter = (today.month - 1) // 3 + 1
        year = today.year

        if preset == "last_quarter":

            quarter -= 1

            if quarter == 0:
                quarter, year = 4, year - 1

        first_month = (quarter - 1) * 3 + 1

        return (
            _month_start(year, first_month),
            _month_end(year, first_month + 2),
            f"{quarter}. štvrťrok {year}",
        )

    if preset == "this_year":
        return date(today.year, 1, 1), date(today.year, 12, 31), str(today.year)

    if preset == "last_year":
        return date(today.year - 1, 1, 1), date(today.year - 1, 12, 31), str(today.year - 1)

    raise ValueError("Neznáme obdobie.")


# =========================================
# TEXTY (ochrana pred vzorcami)
# =========================================

def safe_text(value) -> str:
    """Text bezpečný pre CSV/Excel: začiatok =,+,-,@, tabulátor alebo
    CR by tabuľkový procesor vyhodnotil ako vzorec."""

    if value is None:
        return ""

    text = str(value)

    if text and text[0] in ("=", "+", "-", "@", "\t", "\r"):
        return "'" + text

    return text


# =========================================
# RIADKY
# =========================================

def invoice_type(invoice) -> str:

    if invoice.is_credit_note:
        return "Dobropis"

    if invoice.is_proforma:
        return "Zálohová faktúra"

    return "Faktúra"


def select_invoices(invoices, start: date, end: date, include_proforma: bool = False) -> list:

    selected = [
        invoice for invoice in invoices
        if start <= invoice.issue_date <= end
        and invoice.status in EXPORT_STATUSES
        and (include_proforma or not invoice.is_proforma)
    ]

    selected.sort(key=lambda invoice: (invoice.issue_date, invoice.invoice_number))

    return selected


def invoice_row(invoice, original_numbers: dict[int, str] | None = None) -> dict:

    totals = calculate_invoice_totals(invoice.items)
    sign = -1 if invoice.is_credit_note else 1
    customer = invoice.customer

    original = ""

    if invoice.is_credit_note and invoice.original_invoice_id and original_numbers:
        original = original_numbers.get(invoice.original_invoice_id, "")

    counted = invoice.status != "Stornovaná"

    return {
        "Číslo": invoice.invoice_number,
        "Typ": invoice_type(invoice),
        "Dátum vystavenia": invoice.issue_date,
        "Dátum dodania": invoice.delivery_date,
        "Dátum splatnosti": invoice.due_date,
        "Dátum úhrady": invoice.paid_date,
        "Odberateľ": customer.name if customer else "",
        "IČO": customer.ico if customer else "",
        "DIČ": customer.dic if customer else "",
        "IČ DPH": customer.ic_dph if customer else "",
        "Základ": sign * totals["total_base"],
        "DPH": sign * totals["total_vat"],
        "Spolu": sign * totals["total_gross"],
        "Stav": invoice.status,
        "Spôsob úhrady": invoice.payment_method,
        "Variabilný symbol": invoice.variable_symbol or invoice.invoice_number,
        "Prenesená daňová povinnosť": "Áno" if invoice.reverse_charge else "Nie",
        "K faktúre": original,
        "V súčte": "Áno" if counted else "Nie",
    }


def invoice_rows(invoices, original_numbers: dict[int, str] | None = None) -> list[dict]:
    """Riadky pre export. `original_numbers` (id -> číslo) slúži na
    stĺpec "K faktúre" pri dobropisoch, ktorých pôvodná faktúra môže byť
    mimo obdobia."""

    numbers = {invoice.id: invoice.invoice_number for invoice in invoices}
    numbers.update(original_numbers or {})

    return [invoice_row(invoice, numbers) for invoice in invoices]


def vat_summary(invoices) -> list[dict]:
    """Rekapitulácia DPH podľa sadzieb (bez stornovaných, dobropisy záporne)."""

    by_rate: dict[int, dict] = {}

    for invoice in invoices:

        if invoice.status == "Stornovaná":
            continue

        sign = -1 if invoice.is_credit_note else 1

        for entry in calculate_invoice_totals(invoice.items)["vat_breakdown"]:

            bucket = by_rate.setdefault(entry["rate"], {"rate": entry["rate"], "base": ZERO, "vat": ZERO, "gross": ZERO})

            bucket["base"] += sign * entry["base"]
            bucket["vat"] += sign * entry["vat"]
            bucket["gross"] += sign * entry["gross"]

    return sorted(by_rate.values(), key=lambda entry: -entry["rate"])


def cost_rows(costs, start: date, end: date) -> list[dict]:

    rows = []

    for cost in costs:

        if not (start <= cost.cost_date <= end):
            continue

        job = cost.job

        rows.append({
            "Dátum": cost.cost_date,
            "Zákazka": job.title if job else "",
            "Zákazník": job.customer.name if job and job.customer else "",
            "Druh": COST_LABELS.get(cost.category, "Iné"),
            "Popis": cost.description,
            "Hodiny": cost.hours,
            "Suma bez DPH": cost.amount,
        })

    rows.sort(key=lambda row: (row["Dátum"], row["Zákazka"]))

    return rows


def totals_of(rows: list[dict]) -> dict:
    """Súčty faktúr zahrnutých do súčtu (bez stornovaných)."""

    counted = [row for row in rows if row["V súčte"] == "Áno"]

    return {
        "count": len(rows),
        "counted": len(counted),
        "base": sum((row["Základ"] for row in counted), ZERO),
        "vat": sum((row["DPH"] for row in counted), ZERO),
        "gross": sum((row["Spolu"] for row in counted), ZERO),
        "unpaid": sum((row["Spolu"] for row in counted if row["Stav"] == "Odoslaná"), ZERO),
    }


# =========================================
# CSV
# =========================================

def _csv_value(value) -> str:

    if value is None:
        return ""

    if isinstance(value, date):
        return value.strftime("%d.%m.%Y")

    if isinstance(value, Decimal):
        return f"{value:.2f}".replace(".", ",")

    return safe_text(value)


def build_csv(rows: list[dict], columns: tuple) -> bytes:
    """CSV pre slovenský Excel: UTF-8 s BOM, oddeľovač ';', desatinná čiarka."""

    buffer = io.StringIO()

    writer = csv.writer(buffer, delimiter=";", lineterminator="\r\n")

    writer.writerow(columns)

    for row in rows:
        writer.writerow([_csv_value(row.get(column)) for column in columns])

    return buffer.getvalue().encode("utf-8-sig")


# =========================================
# EXCEL
# =========================================

HEADER_FONT = Font(bold=True, color="FFFFFF")
HEADER_FILL = PatternFill("solid", fgColor="2F3E46")
MONEY_FORMAT = "#,##0.00"
DATE_FORMAT = "DD.MM.YYYY"


def _write_table(sheet, columns: tuple, rows: list[dict], money: set[str], dates: set[str]) -> None:

    sheet.append(list(columns))

    for cell in sheet[1]:

        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(vertical="center", wrap_text=True)

    for row in rows:

        values = []

        for column in columns:

            value = row.get(column)

            if isinstance(value, Decimal):
                value = float(value)

            elif isinstance(value, str):
                value = safe_text(value)

            values.append(value)

        sheet.append(values)

    for index, column in enumerate(columns, start=1):

        letter = get_column_letter(index)

        longest = max([len(str(column))] + [len(str(row.get(column) or "")) for row in rows[:200]])

        sheet.column_dimensions[letter].width = min(max(12, longest + 2), 48)

        for cell in sheet[letter][1:]:

            if column in money:
                cell.number_format = MONEY_FORMAT

            elif column in dates:
                cell.number_format = DATE_FORMAT

    sheet.freeze_panes = "A2"

    if rows:
        sheet.auto_filter.ref = sheet.dimensions


def build_xlsx(*, period_label: str, company, invoices_rows: list[dict], vat_rows: list[dict],
               costs: list[dict], generated: date) -> bytes:

    totals = totals_of(invoices_rows)

    costs_total = sum((row["Suma bez DPH"] for row in costs), ZERO)

    workbook = Workbook()

    # ---- Súhrn ----
    summary = workbook.active
    summary.title = "Súhrn"

    lines = [
        ("Export pre účtovníka", None),
        ("Obdobie", period_label),
        ("Firma", safe_text(company.name) if company and company.name else ""),
        ("IČO", safe_text(company.ico) if company and company.ico else ""),
        ("DIČ", safe_text(company.dic) if company and company.dic else ""),
        ("IČ DPH", safe_text(company.ic_dph) if company and company.ic_dph else ""),
        ("Vygenerované", generated),
        (None, None),
        ("Počet dokladov (bez stornovaných)", totals["counted"]),
        ("Počet stornovaných", totals["count"] - totals["counted"]),
        ("Základ dane", float(totals["base"])),
        ("DPH", float(totals["vat"])),
        ("Spolu s DPH", float(totals["gross"])),
        ("Z toho neuhradené", float(totals["unpaid"])),
        (None, None),
        ("Náklady zo zákaziek (bez DPH)", float(costs_total)),
        ("Základ − náklady", float(totals["base"] - costs_total)),
    ]

    for label, value in lines:
        summary.append([label, value])

    summary["A1"].font = Font(bold=True, size=14)

    for row in summary.iter_rows(min_row=2):

        row[0].font = Font(bold=True)

        if isinstance(row[1].value, float):
            row[1].number_format = MONEY_FORMAT

        elif isinstance(row[1].value, date):
            row[1].number_format = DATE_FORMAT

        row[1].alignment = Alignment(horizontal="left")

    summary.column_dimensions["A"].width = 38
    summary.column_dimensions["B"].width = 30

    # ---- Faktúry ----
    sheet = workbook.create_sheet("Faktúry")

    _write_table(
        sheet, INVOICE_COLUMNS, invoices_rows,
        money={"Základ", "DPH", "Spolu"},
        dates={"Dátum vystavenia", "Dátum dodania", "Dátum splatnosti", "Dátum úhrady"},
    )

    # ---- DPH ----
    vat_sheet = workbook.create_sheet("Rekapitulácia DPH")

    vat_sheet.append(["Sadzba DPH", "Základ", "DPH", "Spolu"])

    for cell in vat_sheet[1]:

        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL

    for entry in vat_rows:
        vat_sheet.append([f"{entry['rate']} %", float(entry["base"]), float(entry["vat"]), float(entry["gross"])])

    vat_sheet.append(["Spolu", float(totals["base"]), float(totals["vat"]), float(totals["gross"])])

    for cell in vat_sheet[vat_sheet.max_row]:
        cell.font = Font(bold=True)

    for letter in "BCD":

        vat_sheet.column_dimensions[letter].width = 16

        for cell in vat_sheet[letter][1:]:
            cell.number_format = MONEY_FORMAT

    vat_sheet.column_dimensions["A"].width = 14

    # ---- Náklady ----
    cost_sheet = workbook.create_sheet("Náklady")

    _write_table(
        cost_sheet, COST_COLUMNS, costs,
        money={"Suma bez DPH"}, dates={"Dátum"},
    )

    buffer = io.BytesIO()
    workbook.save(buffer)

    return buffer.getvalue()
